# -*- coding: utf-8 -*-
"""提取参考文档（AzureDoc）的排版参数，用于对齐我们的 Word 文档。"""

import docx
from docx.shared import Pt

REF = r"F:\AzureCore\AzureDoc\About_EXP2_System_CN.docx"
doc = docx.Document(REF)

print("=== 页面设置 ===")
s = doc.sections[0]
print(f"页面: {s.page_width.cm:.1f} x {s.page_height.cm:.1f} cm | 边距 上{s.top_margin.cm:.1f} 下{s.bottom_margin.cm:.1f} 左{s.left_margin.cm:.1f} 右{s.right_margin.cm:.1f}")

print("\n=== 前 14 段的样式与字号 ===")
for p in doc.paragraphs[:14]:
    if not p.text.strip():
        continue
    runs = [r for r in p.runs if r.text.strip()]
    sizes = {r.font.size.pt for r in runs if r.font.size}
    fonts = {r.font.name for r in runs if r.font.name}
    bold = any(r.bold for r in runs)
    print(f"[{p.style.name:12s}] align={p.alignment} size={sizes or '继承'} font={fonts or '继承'} bold={bold} | {p.text[:56]!r}")

print("\n=== 各级标题的字号 ===")
seen = {}
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") and p.text.strip():
        runs = [r for r in p.runs if r.text.strip()]
        sizes = [r.font.size.pt for r in runs if r.font.size]
        fonts = [r.font.name for r in runs if r.font.name]
        key = p.style.name
        if key not in seen:
            seen[key] = (sizes, fonts, p.text[:40])
for k, v in seen.items():
    print(f"{k}: size={v[0]} font={v[1]} | 例: {v[2]!r}")

print("\n=== 正文段落字号统计 ===")
from collections import Counter
c = Counter()
for p in doc.paragraphs:
    if p.style.name == "Normal" and p.text.strip():
        for r in p.runs:
            if r.font.size:
                c[r.font.size.pt] += 1
print(dict(c))

print("\n=== 表格样式 ===")
for t in doc.tables:
    print(f"{len(t.rows)}x{len(t.columns)} style={t.style.name if t.style else '无'}")
