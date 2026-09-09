# -*- coding: utf-8 -*-
"""诊断生成的 docx 结构问题。"""

from collections import Counter

import docx

doc = docx.Document(r"E:\IUC\risk-warning-platform\docs\设计哲学-企业经营风险预警平台.docx")
texts = [p.text for p in doc.paragraphs]
print("段落总数:", len(texts))
print("含 ** 残留:", sum(1 for t in texts if "**" in t))
print("含 | 残留:", sum(1 for t in texts if t.strip().startswith("|")))
print("含 --- 残留:", sum(1 for t in texts if t.strip() == "---"))
print("含 ` 残留:", sum(1 for t in texts if "`" in t))
print()
print("--- 样式统计 ---")
for k, v in Counter(p.style.name for p in doc.paragraphs).items():
    print(f"  {k}: {v}")
print()
print("--- 封面后 8 段 ---")
for p in doc.paragraphs[:8]:
    print(f"  [{p.style.name}] {p.text[:70]!r}")
print()
print("--- 3.1 架构图段落 ---")
for p in doc.paragraphs:
    if "表现层" in p.text:
        print("  文本:", repr(p.text[:160]))
        print("  runs:", len(p.runs), "| style:", p.style.name)
        break
print()
print("--- 表格样式 ---")
for t in doc.tables:
    print("  表格:", len(t.rows), "x", len(t.columns), "| style:", t.style.name if t.style else "无")
