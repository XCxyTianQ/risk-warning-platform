# -*- coding: utf-8 -*-
"""把设计哲学 Markdown 转成 Word（.docx），排版严格对齐 AzureDoc 既有技术文档。

对齐参数（取自 About_EXP2_System_CN.docx）：
- 页面：A4，四周 2.5cm 边距
- 封面：主标题 28pt 加粗居中 / 副标题 14pt 居中 / 项目与日期 12pt 居中 / 定位 11pt 居中
- 正文：12pt，行距 1.5
- 标题：`##` → Heading 1，`###` → Heading 2（保证 Word 大纲有正确的一级标题）
- 表格：Table Grid，表头加粗 + 浅底纹
- 代码块：Consolas 10pt + 浅灰底纹

用法：python docs/scripts/md_to_docx.py
"""

import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "设计哲学-企业经营风险预警平台.md"
DST = ROOT / "docs" / "设计哲学-企业经营风险预警平台.docx"

CJK = "微软雅黑"
LATIN = "Times New Roman"

doc = docx.Document()

# ---------- 页面 ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)


def set_font(run, size=None, bold=None, italic=None, cjk=CJK, latin=LATIN, mono=False):
    if mono:
        cjk = latin = "Consolas"
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def shade(par, color="F2F2F2"):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


# ---------- 样式基线 ----------
normal = doc.styles["Normal"]
normal.font.size = Pt(12)
normal.font.name = LATIN
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

for lvl, size in ((1, 16), (2, 14), (3, 12.5)):
    st = doc.styles[f"Heading {lvl}"]
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.name = LATIN
    st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), CJK)
    rf.set(qn("w:ascii"), LATIN)
    rf.set(qn("w:hAnsi"), LATIN)
    st.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
    st.paragraph_format.space_after = Pt(6)


def add_rich(par, text: str, size=12) -> None:
    """处理 **加粗** 与 `代码` 行内标记。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            set_font(run, size=size - 0.5, mono=True)
        else:
            run = par.add_run(part)
            set_font(run, size=size)


def cover(title: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(title), size=28, bold=True)

    if lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        set_font(p.add_run(lines[0]), size=14)

    rest = lines[1:]
    if rest:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        set_font(p.add_run("\n".join(rest[:2])), size=12)

    tail = rest[2:]
    if tail:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(24)
        set_font(p.add_run("\n".join(tail)), size=11, italic=True, cjk="楷体")

    doc.add_page_break()


lines = SRC.read_text(encoding="utf-8").splitlines()
i = 0
cover_done = False
pending_cover: list[str] = []

while i < len(lines):
    raw = lines[i].rstrip()
    line = raw.strip()

    # 封面：一级标题 + 其后的副标题/项目/日期/定位（直到第一个二级标题）
    if not cover_done:
        m = re.match(r"^#\s+(.*)$", raw)
        if m:
            title = m.group(1).strip()
            j = i + 1
            while j < len(lines):
                nxt = lines[j].strip()
                if nxt.startswith("##"):
                    break
                if nxt and nxt not in ("---", "***"):
                    pending_cover.append(nxt.lstrip("> ").strip().replace("**", ""))
                j += 1
            cover(title, [x for x in pending_cover if x])
            cover_done = True
            i = j
            continue
        i += 1
        continue

    # 代码块
    if line.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(10)
        set_font(p.add_run("\n".join(buf)), size=10, mono=True)
        shade(p)
        continue

    # 表格
    if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
        header = [c.strip() for c in line.strip("|").split("|")]
        i += 2
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ci, h in enumerate(header):
            cell = table.rows[0].cells[ci]
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(h), size=11, bold=True)
            shade(cell.paragraphs[0], "E7EEF8")
        for row in rows:
            cells = table.add_row().cells
            for ci, val in enumerate(row[: len(header)]):
                cells[ci].text = ""
                add_rich(cells[ci].paragraphs[0], val, size=11)
        doc.add_paragraph()
        continue

    # 标题
    m = re.match(r"^(#{2,4})\s+(.*)$", raw)
    if m:
        level = len(m.group(1)) - 1  # ## → Heading 1, ### → Heading 2
        doc.add_heading(m.group(2).strip(), level=level)
        i += 1
        continue

    # 引用
    if line.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(8)
        add_rich(p, line.lstrip("> ").strip(), size=11)
        for run in p.runs:
            run.font.italic = True
        i += 1
        continue

    # 分隔线
    if line in ("---", "***"):
        i += 1
        continue

    # 列表
    if re.match(r"^\s*[-*]\s+", raw) or re.match(r"^\s*\d+\.\s+", raw):
        style = "List Bullet" if line[0] in "-*" else "List Number"
        p = doc.add_paragraph(style=style)
        add_rich(p, re.sub(r"^\s*([-*]|\d+\.)\s+", "", raw), size=12)
        i += 1
        continue

    if not line:
        i += 1
        continue

    p = doc.add_paragraph()
    add_rich(p, raw, size=12)
    i += 1

doc.save(DST)
print("saved:", DST)
